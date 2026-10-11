import csv
import json
import os
import re
import sys
from datetime import datetime
from pathlib import Path


MAX_SCORE = 10
SCORE_BY_SOLVED = {0: 1, 1: 3, 2: 5, 3: 7, 4: 9, 5: 10}   


def score_for(solved):
    
    return SCORE_BY_SOLVED[max(0, min(int(solved), max(SCORE_BY_SOLVED)))]


def safe_name(value):
    return re.sub(r"[^\w\-]+", "_", str(value), flags=re.U).strip("_") or "x"


def _nat_key(text):
    
    return [int(p) if p.isdigit() else p.casefold() for p in re.split(r"(\d+)", str(text))]



_MONO = ["consola.ttf", "C:/Windows/Fonts/consola.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf",
         "DejaVuSansMono.ttf", "/usr/share/fonts/truetype/liberation/LiberationMono-Regular.ttf",
         "/System/Library/Fonts/Menlo.ttc", "Menlo.ttc", "cour.ttf"]
_UI = ["segoeui.ttf", "C:/Windows/Fonts/segoeui.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
       "DejaVuSans.ttf", "/System/Library/Fonts/Supplemental/Arial.ttf", "arial.ttf"]
_UI_BOLD = ["segoeuib.ttf", "C:/Windows/Fonts/segoeuib.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
            "DejaVuSans-Bold.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf", "arialbd.ttf"]

_COL = dict(
    bg="#0c0e1d", card="#1a1e38", card_edge="#2b3058", chip="#252a50",
    code_bg="#0e1020", fg="#e9ebfa", muted="#8b91b8", accent2="#22d3ee",
    ok="#34d399", ok_text="#07150f", ok_card="#0f3b2c", ok_code="#0b261d",
    bad="#fb7185", warn="#fbbf24",
)

def _font(candidates, size):
    from PIL import ImageFont
    for path in candidates:
        try:
            return ImageFont.truetype(path, size)
        except OSError:
            continue
    try:
        return ImageFont.load_default(size)
    except TypeError:
        return ImageFont.load_default()


def _fit(candidates, text, size, max_w, min_size=18):
    
    while size > min_size:
        f = _font(candidates, size)
        if f.getlength(text) <= max_w:
            return f
        size -= 2
    return _font(candidates, min_size)


def _layout_code(code, cols):
    
    code = (code or "").replace("\r\n", "\n").replace("\t", "    ").rstrip()
    if not code.strip():
        return [(None, "(решение не написано)")], True
    rows = []
    for n, line in enumerate(code.split("\n"), 1):
        line = line.rstrip()
        chunks = [line[i:i + cols] for i in range(0, len(line), cols)] or [""]
        for k, ch in enumerate(chunks):
            rows.append((n if k == 0 else None, ch))
    return rows, False


def build_solutions_image(path, name, group, items, solved, score, reason="", when=None):
    
    from PIL import Image, ImageDraw

    when = when or datetime.now()
    W, M, PAD = 1500, 30, 22
    f_code, f_num = _font(_MONO, 20), _font(_MONO, 18)
    f_h2, f_txt = _font(_UI_BOLD, 27), _font(_UI, 21)
    f_small, f_badge = _font(_UI, 18), _font(_UI_BOLD, 19)
    f_big = _font(_UI_BOLD, 66)
    f_name = _fit(_UI_BOLD, name, 44, W - 2 * M - 400)
    line_h = 28
    char_w = f_code.getlength("M")
    gutter = int(char_w * 3 + 22)
    code_inner_w = W - 2 * M - 2 * PAD - 2 * 14 - gutter
    cols = max(30, int(code_inner_w // char_w))

    panels = []
    for it in items:
        rows, empty = _layout_code(it.get("code", ""), cols)
        code_h = len(rows) * line_h + 28
        condition = None
        image_path = str(it.get("image", "") or "").strip()
        if image_path:
            try:
                with Image.open(image_path) as source:
                    condition = source.convert("RGB")
            except (OSError, ValueError) as e:
                raise OSError(f"Не удалось открыть изображение условия «{image_path}»: {e}") from e
            condition.thumbnail((code_inner_w, 420))
        condition_h = condition.height + 46 if condition else 0
        panels.append(dict(it=it, rows=rows, empty=empty, code_h=code_h, condition=condition,
                           h=PAD + 62 + condition_h + code_h + PAD))

    banner_h, footer_h, gap = 210, 76, 22
    H = banner_h + sum(p["h"] + gap for p in panels) + footer_h
    img = Image.new("RGB", (W, H), _COL["bg"])
    d = ImageDraw.Draw(img)

    
    for x in range(W):                                   
        t = x / (W - 1)
        c1, c2 = (124, 92, 255), (34, 211, 238)
        d.line([(x, 0), (x, 8)], fill=tuple(round(a + (b - a) * t) for a, b in zip(c1, c2)))
    d.text((M, 34), "ФОРМАТИВНАЯ ОЦЕНКА  ·  РЕШЕНИЯ УЧЕНИКА", font=f_small, fill=_COL["muted"])
    d.text((M, 64), name, font=f_name, fill=_COL["fg"])
    d.text((M, 126), f"Группа: {group}", font=f_txt, fill=_COL["accent2"])
    d.text((M, 160), f"{when:%d.%m.%Y %H:%M}" + (f"   ·   {reason}" if reason else ""), font=f_small, fill=_COL["muted"])
    bx0, by0, bx1, by1 = W - M - 330, 30, W - M, 180
    d.rounded_rectangle((bx0, by0, bx1, by1), radius=22, fill=_COL["card"], outline=_COL["card_edge"], width=2)
    d.text((bx0 + 24, by0 + 16), "ОЦЕНКА", font=f_small, fill=_COL["muted"])
    d.text((bx0 + 24, by0 + 38), f"{score}", font=f_big, fill=_COL["ok"] if score >= 7 else (_COL["warn"] if score >= 5 else _COL["bad"]))
    d.text((bx0 + 24 + f_big.getlength(f"{score}") + 8, by0 + 76), f"/ {MAX_SCORE}", font=f_h2, fill=_COL["muted"])
    d.text((bx0 + 24, by1 - 38), f"решено {solved} из {len(items)}", font=f_txt, fill=_COL["fg"])

    
    y = banner_h
    for idx, p in enumerate(panels, 1):
        it = p["it"]
        ok = bool(it.get("solved"))
        fill, edge = (_COL["ok_card"], _COL["ok"]) if ok else (_COL["card"], _COL["card_edge"])
        d.rounded_rectangle((M, y, W - M, y + p["h"]), radius=20, fill=fill, outline=edge, width=4 if ok else 2)
        cx, cy = M + PAD + 24, y + PAD + 26
        d.ellipse((cx - 24, cy - 24, cx + 24, cy + 24), fill=_COL["ok"] if ok else _COL["chip"],
                  outline=edge, width=2)
        if ok:                                           
            d.line([(cx - 11, cy + 1), (cx - 3, cy + 10), (cx + 12, cy - 9)], fill=_COL["ok_text"], width=5, joint="curve")
        else:
            nw = f_h2.getlength(str(idx))
            d.text((cx - nw / 2, cy - 17), str(idx), font=f_h2, fill=_COL["fg"])
        title = f"Задача {idx} · {it.get('title', '')}"
        d.text((cx + 40, y + PAD - 2), title, font=f_h2, fill=_COL["fg"])
        meta = f"{it.get('lang', '')}   ·   баллы {it.get('score', 0)} / {it.get('max_points', 0)}"
        if it.get("total"):
            meta += f"   ·   тестов {it.get('passed', 0)} / {it.get('total', 0)}"
        d.text((cx + 40, y + PAD + 34), meta, font=f_small, fill=_COL["accent2"])

        label = "РЕШЕНО" if ok else "НЕ РЕШЕНО"
        lw = f_badge.getlength(label) + 40
        bx1 = W - M - PAD
        bx0 = bx1 - lw
        d.rounded_rectangle((bx0, y + PAD + 6, bx1, y + PAD + 46), radius=20,
                            fill=_COL["ok"] if ok else None, outline=None if ok else _COL["card_edge"], width=2)
        d.text((bx0 + 20, y + PAD + 15), label, font=f_badge, fill=_COL["ok_text"] if ok else _COL["muted"])

        
        cy0 = y + PAD + 62
        if p["condition"]:
            d.text((M + PAD + 14, cy0 + 2), "УСЛОВИЕ", font=f_small, fill=_COL["muted"])
            image_y = cy0 + 28
            img.paste(p["condition"], (M + PAD + 14, image_y))
            cy0 = image_y + p["condition"].height + 18
        d.rounded_rectangle((M + PAD, cy0, W - M - PAD, cy0 + p["code_h"]), radius=14,
                            fill=_COL["ok_code"] if ok else _COL["code_bg"])
        ty = cy0 + 14
        tx = M + PAD + 14 + gutter
        for num, text in p["rows"]:
            if num is not None:
                s = str(num)
                d.text((tx - 14 - f_num.getlength(s), ty + 2), s, font=f_num, fill=_COL["muted"])
            stripped = text.lstrip()
            dim = p["empty"] or stripped.startswith(("#", "//", "/*"))
            d.text((tx, ty), text, font=f_code, fill=_COL["muted"] if dim else _COL["fg"])
            ty += line_h
        y += p["h"] + gap

    d.text((M, H - footer_h + 20), f"Подпись: {name} · {group} · {when:%d.%m.%Y %H:%M:%S}",
           font=f_small, fill=_COL["muted"])
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    img.save(path, "PNG", optimize=True)
    return Path(path)



def save_submission(folder, base, name, group, items, reason=""):
    

    
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    when = datetime.now()
    solved = sum(1 for it in items if it.get("solved"))
    score = score_for(solved)
    out = {"solved": solved, "score": score, "png": None, "json": None, "csv": None, "errors": []}

    png_path = folder / f"{safe_name(name)}.png"
    try:
        build_solutions_image(png_path, name, group, items, solved, score, reason, when)
        out["png"] = png_path
    except ImportError:
        out["errors"].append("не установлен Pillow (pip install pillow): картинка с решениями не создана")
        return out
    except Exception as e:
        out["errors"].append(f"картинка не создана: {e}")
        return out

    data = {
        "version": 1, "name": name, "group": group, "finished": when.isoformat(timespec="seconds"),
        "reason": reason, "solved_count": solved, "total_tasks": len(items), "score": score, "max_score": MAX_SCORE,
        "tasks": [{"title": it.get("title", ""), "solved": bool(it.get("solved")), "lang": it.get("lang", ""),
                   "score": it.get("score", 0), "max_points": it.get("max_points", 0)} for it in items],
    }
    json_path = folder / f"result_{base}.json"
    tmp = json_path.with_suffix(".tmp")          
    tmp.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(tmp, json_path)
    out["json"] = json_path

    try:
        rows, _ = collect_results(folder)
        out["csv"] = write_csv(rows, folder)
    except OSError as e:
        out["errors"].append(f"сводки CSV не обновлены: {e}")
    return out



def collect_results(folder):
    
    rows, problems, latest = [], [], {}
    for p in sorted(Path(folder).glob("result_*.json")):
        try:
            d = json.loads(p.read_text(encoding="utf-8"))
            name, group = str(d["name"]).strip(), str(d["group"]).strip()
            solved = int(d["solved_count"])
            tasks = [bool(t.get("solved")) for t in d.get("tasks", [])]
            total = int(d.get("total_tasks", len(tasks)))
            finished = str(d.get("finished", ""))
        except (OSError, ValueError, KeyError, TypeError, AttributeError) as e:
            problems.append(f"{p.name}: {e}")
            continue
        row = dict(name=name, group=group, solved=solved, total=max(total, solved), tasks=tasks,
                   score=score_for(solved), finished=finished, file=p.name)
        key = (group.casefold(), name.casefold())
        if key not in latest or finished >= latest[key]["finished"]:
            latest[key] = row
    rows = sorted(latest.values(), key=lambda r: (_nat_key(r["group"]), r["name"].casefold()))
    return rows, problems


def group_stats(rows):
    stats = {}
    for r in rows:
        stats.setdefault(r["group"], []).append(r)
    out = []
    for g in sorted(stats, key=_nat_key):
        sc = [r["score"] for r in stats[g]]
        out.append(dict(group=g, count=len(sc), avg=sum(sc) / len(sc), best=max(sc), worst=min(sc),
                        perfect=sum(1 for r in stats[g] if r["solved"] >= r["total"] > 0)))
    return out


def _comma(x):
    return f"{x:.2f}".replace(".", ",")     


def write_csv(rows, folder):
   
    folder = Path(folder)
    p1, p2 = folder / "results.csv", folder / "groups.csv"
    with open(p1, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Группа", "ФИО", "Решено задач", f"Балл (из {MAX_SCORE})"])
        for r in rows:
            w.writerow([r["group"], r["name"], r["solved"], r["score"]])
    with open(p2, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Группа", "Участников", "Средний балл", "Лучший балл", "Худший балл", "Решили все задачи"])
        for s in group_stats(rows):
            w.writerow([s["group"], s["count"], _comma(s["avg"]), s["best"], s["worst"], s["perfect"]])
    return p1, p2


def write_xlsx(rows, folder):
    
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    head_fill, head_font = PatternFill("solid", fgColor="1F2347"), Font(bold=True, color="FFFFFF")
    ok_fill, no_fill = PatternFill("solid", fgColor="C6F6D5"), PatternFill("solid", fgColor="EEF0F8")
    score_fill = {1: "FECACA", 3: "FED7AA", 5: "FDE68A", 7: "D9F99D", 9: "A7F3D0", 10: "6EE7B7"}
    thin = Side(style="thin", color="C9CDE6")
    center = Alignment(horizontal="center", vertical="center")

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill, cell.font, cell.alignment = head_fill, head_font, center
        ws.row_dimensions[1].height = 26
        ws.freeze_panes = "A2"

    n_tasks = max([len(r["tasks"]) for r in rows] + [max(r["total"] for r in rows) if rows else 1, 1])
    wb = Workbook()

    ws = wb.active
    ws.title = "Общий список"
    ws.append(["№", "Группа", "ФИО"] + [f"Задача {i}" for i in range(1, n_tasks + 1)] + ["Решено", f"Балл (из {MAX_SCORE})"])
    prev_group = None
    for i, r in enumerate(rows, 1):
        flags = (r["tasks"] + [False] * n_tasks)[:n_tasks]
        ws.append([i, r["group"], r["name"]] + ["✔" if f else "—" for f in flags] + [r["solved"], r["score"]])
        row_idx = ws.max_row
        for k, f in enumerate(flags):
            cell = ws.cell(row=row_idx, column=4 + k)
            cell.fill, cell.alignment = (ok_fill if f else no_fill), center
        sc = ws.cell(row=row_idx, column=4 + n_tasks + 1)
        sc.fill = PatternFill("solid", fgColor=score_fill.get(r["score"], "FFFFFF"))
        sc.font, sc.alignment = Font(bold=True), center
        ws.cell(row=row_idx, column=4 + n_tasks).alignment = center
        ws.cell(row=row_idx, column=1).alignment = center
        if prev_group is not None and r["group"] != prev_group:      
            for c in range(1, 4 + n_tasks + 2):
                ws.cell(row=row_idx, column=c).border = Border(top=Side(style="medium", color="7C5CFF"))
        prev_group = r["group"]
    style_header(ws, 3 + n_tasks + 2)
    ws.auto_filter.ref = ws.dimensions
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    widths = [5, 12, 34] + [10] * n_tasks + [10, 16]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws2 = wb.create_sheet("По группам")
    ws2.append(["Группа", "Участников", "Средний балл", "Лучший балл", "Худший балл", "Решили все задачи"])
    stats = group_stats(rows)
    for s in stats:
        ws2.append([s["group"], s["count"], round(s["avg"], 2), s["best"], s["worst"], s["perfect"]])
        ws2.cell(row=ws2.max_row, column=3).number_format = "0.00"
    if rows:
        all_sc = [r["score"] for r in rows]
        ws2.append(["ВСЕГО", len(rows), round(sum(all_sc) / len(all_sc), 2), max(all_sc), min(all_sc),
                    sum(s["perfect"] for s in stats)])
        last = ws2.max_row
        ws2.cell(row=last, column=3).number_format = "0.00"
        for c in range(1, 7):
            ws2.cell(row=last, column=c).font = Font(bold=True)
            ws2.cell(row=last, column=c).border = Border(top=Side(style="medium", color="7C5CFF"))
    style_header(ws2, 6)
    for i, w in enumerate([14, 13, 14, 14, 14, 18], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    for row in ws2.iter_rows(min_row=2):
        for cell in row[1:]:
            cell.alignment = center

    ws3 = wb.create_sheet("Шкала")
    ws3.append(["Решено задач", f"Балл (из {MAX_SCORE})"])
    for k, v in SCORE_BY_SOLVED.items():
        ws3.append([k, v])
        ws3.cell(row=ws3.max_row, column=2).fill = PatternFill("solid", fgColor=score_fill.get(v, "FFFFFF"))
    style_header(ws3, 2)
    ws3.column_dimensions["A"].width, ws3.column_dimensions["B"].width = 16, 16
    for row in ws3.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = center
    ws3["D1"] = "Задача считается решённой, если пройдены все тесты."

    path = Path(folder) / "results.xlsx"
    wb.save(path)
    return path


def write_summaries(rows, folder):
   
    files, warnings = [], []
    files += list(write_csv(rows, folder))
    try:
        files.append(write_xlsx(rows, folder))
    except ImportError:
        if getattr(sys, "frozen", False):
            warnings.append("Excel-файл не создан: openpyxl отсутствует в сборке приложения. "
                            "Требуется пересобрать Teacher.exe. CSV готовы.")
        else:
            warnings.append("Excel-файл не создан: установите openpyxl (pip install openpyxl). CSV готовы.")
    except PermissionError:
        warnings.append("results.xlsx не перезаписан: закройте его в Excel и повторите. CSV готовы.")
    return files, warnings


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path("reports")
    data_rows, probs = collect_results(target)
    made, warns = write_summaries(data_rows, target)
    print(f"Учеников: {len(data_rows)}")
    for f in made:
        print("  ", f)
    for t in probs + warns:
        print("  !", t)